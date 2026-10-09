#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
销售汇总报告生成器。

读取指定目录下所有 xlsx 销售明细（列：日期,门店,商品,数量,单价,金额），
合并后生成《销售汇总报告.xlsx》，包含三个 sheet：
  总览      - 总销售额、订单总数、客单价、覆盖月份
  按商品汇总 - 各商品销售额（降序）
  按门店汇总 - 各门店销售额（降序）

用法:
    python sales_report.py                              # 读 ./data/，输出 ./销售汇总报告.xlsx
    python sales_report.py --input ./data --output report.xlsx
"""

import argparse
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, numbers

EXPECTED_HEADERS = ["日期", "门店", "商品", "数量", "单价", "金额"]
MONEY_FMT = '#,##0.00'


def display_width(text: str) -> int:
    """估算显示宽度：中文算 2，ASCII 算 1。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def auto_width(ws, min_width: int = 10, max_width: int = 30) -> None:
    for col in ws.columns:
        width = min_width
        for cell in col:
            if cell.value is not None:
                width = max(width, display_width(str(cell.value)) + 2)
        ws.column_dimensions[col[0].column_letter].width = min(width, max_width)


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")


def read_details(input_dir: Path) -> tuple[list, list]:
    """读取目录下所有 xlsx，返回 (明细行列表, 成功处理的文件名列表)。"""
    xlsx_files = sorted(input_dir.glob("*.xlsx"))
    if not xlsx_files:
        print(f"[ERROR] 目录 {input_dir}/ 下没有找到 xlsx 文件", file=sys.stderr)
        sys.exit(1)

    rows, ok_files = [], []
    for path in xlsx_files:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:  # 文件损坏 / 不是合法 xlsx
            print(f"[WARN] 跳过无法读取的文件 {path.name}：{e}")
            continue
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if header != EXPECTED_HEADERS:
            print(f"[WARN] 跳过表头不符的文件 {path.name}（期望 {EXPECTED_HEADERS}）")
            continue
        n = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None:  # 跳过空行
                continue
            rows.append({
                "日期": str(r[0])[:10],
                "门店": r[1],
                "商品": r[2],
                "数量": r[3],
                "单价": r[4],
                "金额": r[5],
            })
            n += 1
        ok_files.append(path.name)
        print(f"[OK] {path.name}: {n} 行")
    if not rows:
        print("[ERROR] 没有读到任何有效明细行", file=sys.stderr)
        sys.exit(1)
    return rows, ok_files


def sheet_overview(wb: Workbook, rows: list) -> None:
    ws = wb.active
    ws.title = "总览"
    total = round(sum(r["金额"] for r in rows), 2)
    count = len(rows)
    avg = round(total / count, 2)
    months = sorted({r["日期"][:7] for r in rows})
    stores = {r["门店"] for r in rows}
    products = {r["商品"] for r in rows}
    data = [
        ("指标", "数值"),
        ("总销售额（元）", total),
        ("订单总数", count),
        ("客单价（元）", avg),
        ("覆盖月份", "、".join(months)),
        ("覆盖门店数", len(stores)),
        ("在售商品数", len(products)),
    ]
    for k, v in data:
        ws.append([k, v])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=4, min_col=2, max_col=2):
        row[0].number_format = MONEY_FMT
    auto_width(ws)


def sheet_by_product(wb: Workbook, rows: list) -> None:
    ws = wb.create_sheet("按商品汇总")
    agg = defaultdict(lambda: {"qty": 0, "amount": 0.0})
    for r in rows:
        agg[r["商品"]]["qty"] += r["数量"]
        agg[r["商品"]]["amount"] += r["金额"]
    ws.append(["商品", "销售数量", "销售额（元）"])
    for name, d in sorted(agg.items(), key=lambda kv: kv[1]["amount"], reverse=True):
        ws.append([name, d["qty"], round(d["amount"], 2)])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = MONEY_FMT
    auto_width(ws)


def sheet_by_store(wb: Workbook, rows: list) -> None:
    ws = wb.create_sheet("按门店汇总")
    agg = defaultdict(lambda: {"orders": 0, "amount": 0.0})
    for r in rows:
        agg[r["门店"]]["orders"] += 1
        agg[r["门店"]]["amount"] += r["金额"]
    ws.append(["门店", "订单数", "销售额（元）"])
    for name, d in sorted(agg.items(), key=lambda kv: kv[1]["amount"], reverse=True):
        ws.append([name, d["orders"], round(d["amount"], 2)])
    style_header(ws)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = MONEY_FMT
    auto_width(ws)


def main() -> None:
    parser = argparse.ArgumentParser(description="合并销售明细 xlsx 并生成汇总报告")
    parser.add_argument("--input", default="data", help="明细数据目录（默认 data/）")
    parser.add_argument("--output", default="销售汇总报告.xlsx", help="输出文件名")
    args = parser.parse_args()

    input_dir = Path(args.input)
    if not input_dir.is_dir():
        print(f"[ERROR] 数据目录不存在：{input_dir}/", file=sys.stderr)
        sys.exit(1)

    rows, ok_files = read_details(input_dir)

    wb = Workbook()
    sheet_overview(wb, rows)
    sheet_by_product(wb, rows)
    sheet_by_store(wb, rows)
    wb.save(args.output)

    print(f"合并 {len(ok_files)} 个文件，共 {len(rows)} 行明细")
    print(f"报告已生成：{args.output}（总览 / 按商品汇总 / 按门店汇总）")


if __name__ == "__main__":
    main()
